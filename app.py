from __future__ import annotations

import csv
import hashlib
import hmac
import io
import json
import logging
import os
import random
import re
import secrets
import shutil
import sqlite3
import tempfile
import uuid
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any, Iterable

from fastapi import Cookie, Depends, FastAPI, File, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field

from qcc.config import settings
from qcc.database import connect as database_connect
from qcc.database import database, uses_postgres
from qcc.storage import StorageError, storage
from qcc import access
from qcc.account_schema import initialize_process, sampling_config

APP_NAME = "Quality Sample Randomizer"
APP_VERSION = "2.0.0-quality-command-center"
ALGORITHM_VERSION = "QSR-RANDOM-V1"
BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = Path(os.getenv("QSR_DATA_DIR", str(BASE_DIR / "data"))).resolve()
UPLOAD_DIR = DATA_DIR / "uploads"
EXPORT_DIR = Path(os.getenv("QSR_EXPORT_DIR", str(BASE_DIR / "exports"))).resolve()
BACKUP_DIR = Path(os.getenv("QSR_BACKUP_DIR", str(BASE_DIR / "backups"))).resolve()
LOG_DIR = Path(os.getenv("QSR_LOG_DIR", str(BASE_DIR / "logs"))).resolve()
DB_PATH = DATA_DIR / "quality_randomizer.db"
MAX_FILE_BYTES = settings.max_upload_bytes
MAX_ROWS = settings.max_rows
MAX_COLS = settings.max_cols
SESSION_HOURS = 8
REQUIRED_TABLES = {"users", "accounts", "account_config", "sampling_runs", "sample_records", "audit_events", "sessions", "uploads"}

for d in [DATA_DIR, UPLOAD_DIR, EXPORT_DIR, BACKUP_DIR, LOG_DIR]:
    d.mkdir(parents=True, exist_ok=True)

logging.basicConfig(
    filename=LOG_DIR / "quality_randomizer.log",
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
)
logger = logging.getLogger("qsr")


def utcnow() -> datetime:
    return datetime.now(timezone.utc)


def iso_now() -> str:
    return utcnow().replace(microsecond=0).isoformat()


def connect(db_path: Path = DB_PATH) -> sqlite3.Connection:
    return database_connect(db_path)


@contextmanager
def db() -> Iterable[sqlite3.Connection]:
    with database(DB_PATH) as con:
        yield con


def init_db() -> None:
    if uses_postgres():
        return
    with db() as con:
        con.executescript(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL UNIQUE COLLATE NOCASE,
                password_salt BLOB NOT NULL,
                password_hash BLOB NOT NULL,
                created_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS sessions (
                token_hash TEXT PRIMARY KEY,
                username TEXT NOT NULL,
                created_at TEXT NOT NULL,
                expires_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS accounts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL UNIQUE COLLATE NOCASE,
                active INTEGER NOT NULL DEFAULT 1,
                created_at TEXT NOT NULL,
                updated_at TEXT NOT NULL
            );
            CREATE TABLE IF NOT EXISTS account_config (
                account_id INTEGER PRIMARY KEY,
                coverage_enabled INTEGER NOT NULL DEFAULT 0,
                coverage_period TEXT NOT NULL DEFAULT 'week',
                audits_per_associate INTEGER NOT NULL DEFAULT 3,
                identifier_column_default TEXT,
                associate_column_default TEXT,
                exclude_previously_sampled INTEGER NOT NULL DEFAULT 1,
                case_insensitive_ids INTEGER NOT NULL DEFAULT 1,
                updated_at TEXT NOT NULL,
                FOREIGN KEY(account_id) REFERENCES accounts(id)
            );
            CREATE TABLE IF NOT EXISTS uploads (
                upload_id TEXT PRIMARY KEY,
                original_name TEXT NOT NULL,
                stored_path TEXT NOT NULL,
                file_type TEXT NOT NULL,
                sha256 TEXT NOT NULL,
                size_bytes INTEGER NOT NULL,
                created_at TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'READY'
            );
            CREATE TABLE IF NOT EXISTS sampling_runs (
                run_id TEXT PRIMARY KEY,
                account_id INTEGER NOT NULL,
                upload_id TEXT,
                source_filename TEXT NOT NULL,
                source_hash TEXT NOT NULL,
                sheet_name TEXT NOT NULL,
                header_row INTEGER NOT NULL,
                identifier_column TEXT NOT NULL,
                associate_column TEXT,
                sampling_method TEXT NOT NULL,
                requested_count INTEGER NOT NULL,
                eligible_count INTEGER NOT NULL,
                selected_count INTEGER NOT NULL,
                seed TEXT NOT NULL,
                algorithm_version TEXT NOT NULL,
                coverage_enabled INTEGER NOT NULL DEFAULT 0,
                coverage_period TEXT,
                coverage_quota INTEGER,
                period_key TEXT,
                filters_json TEXT NOT NULL DEFAULT '[]',
                validation_json TEXT NOT NULL DEFAULT '{}',
                coverage_summary_json TEXT NOT NULL DEFAULT '{}',
                app_version TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL,
                created_by TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'COMPLETED',
                void_reason TEXT,
                voided_at TEXT,
                voided_by TEXT,
                FOREIGN KEY(account_id) REFERENCES accounts(id)
            );
            CREATE TABLE IF NOT EXISTS sample_records (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                run_id TEXT NOT NULL,
                identifier_raw TEXT NOT NULL,
                normalized_identifier TEXT NOT NULL,
                associate TEXT,
                selection_sequence INTEGER NOT NULL,
                row_json TEXT NOT NULL,
                period_key TEXT,
                selected_at TEXT NOT NULL,
                FOREIGN KEY(run_id) REFERENCES sampling_runs(run_id)
            );
            CREATE INDEX IF NOT EXISTS ix_sample_identifier ON sample_records(normalized_identifier);
            CREATE INDEX IF NOT EXISTS ix_sample_run ON sample_records(run_id);
            CREATE INDEX IF NOT EXISTS ix_runs_account ON sampling_runs(account_id, created_at DESC);
            CREATE TABLE IF NOT EXISTS audit_events (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                event_type TEXT NOT NULL,
                username TEXT NOT NULL,
                entity_type TEXT,
                entity_id TEXT,
                detail_json TEXT NOT NULL DEFAULT '{}',
                created_at TEXT NOT NULL
            );
            """
        )
        run_cols = {row[1] for row in con.execute("PRAGMA table_info(sampling_runs)").fetchall()}
        if "validation_json" not in run_cols:
            con.execute("ALTER TABLE sampling_runs ADD COLUMN validation_json TEXT NOT NULL DEFAULT '{}'")
        if "coverage_summary_json" not in run_cols:
            con.execute("ALTER TABLE sampling_runs ADD COLUMN coverage_summary_json TEXT NOT NULL DEFAULT '{}'")
        if "app_version" not in run_cols:
            con.execute("ALTER TABLE sampling_runs ADD COLUMN app_version TEXT NOT NULL DEFAULT ''")


init_db()


def audit(event_type: str, username: str, entity_type: str | None = None, entity_id: str | None = None, detail: dict | None = None) -> None:
    with db() as con:
        con.execute(
            "INSERT INTO audit_events(event_type, username, entity_type, entity_id, detail_json, created_at) VALUES(?,?,?,?,?,?)",
            (event_type, username, entity_type, entity_id, json.dumps(detail or {}, ensure_ascii=False), iso_now()),
        )
    logger.info("event=%s user=%s entity=%s:%s", event_type, username, entity_type, entity_id)


def password_hash(password: str, salt: bytes | None = None) -> tuple[bytes, bytes]:
    salt = salt or secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, 420_000)
    return salt, digest


def verify_password(password: str, salt: bytes, expected: bytes) -> bool:
    _, digest = password_hash(password, salt)
    return hmac.compare_digest(digest, expected)


def create_session(username: str, expected_password_hash: bytes | None = None) -> str:
    token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(token.encode()).hexdigest()
    now = utcnow()
    with db() as con:
        if expected_password_hash is not None:
            # Serialize the final credential check and session insertion with
            # password resets and profile deactivation.
            if not uses_postgres():
                con.execute("BEGIN IMMEDIATE")
            sql = """SELECT u.password_hash,p.active FROM users u
                     JOIN user_profiles p ON p.username=u.username WHERE u.username=?"""
            if uses_postgres():
                sql += " FOR UPDATE OF u, p"
            user = con.execute(sql, (username,)).fetchone()
            if not user or not user["active"] or not hmac.compare_digest(user["password_hash"], expected_password_hash):
                raise HTTPException(401, "Invalid username or password")
        con.execute("DELETE FROM sessions WHERE expires_at < ?", (iso_now(),))
        con.execute(
            "INSERT INTO sessions(token_hash, username, created_at, expires_at) VALUES(?,?,?,?)",
            (token_hash, username, now.replace(microsecond=0).isoformat(), (now + timedelta(hours=SESSION_HOURS)).replace(microsecond=0).isoformat()),
        )
    return token


def current_user(qsr_session: str | None = Cookie(default=None)) -> str:
    if not qsr_session:
        raise HTTPException(401, "Authentication required")
    token_hash = hashlib.sha256(qsr_session.encode()).hexdigest()
    with db() as con:
        row = con.execute("""SELECT s.username, s.expires_at, p.active FROM sessions s
                             LEFT JOIN user_profiles p ON p.username=s.username
                             WHERE s.token_hash=?""", (token_hash,)).fetchone()
        if not row:
            raise HTTPException(401, "Invalid session")
        if not row["active"]:
            raise HTTPException(401, "Invalid session")
        expires_at = row["expires_at"]
        if isinstance(expires_at, str):
            expires_at = datetime.fromisoformat(expires_at)
        if expires_at < utcnow():
            con.execute("DELETE FROM sessions WHERE token_hash=?", (token_hash,))
            raise HTTPException(401, "Session expired")
        return row["username"]


def ensure_roles(username: str, *allowed: str) -> None:
    with db() as con:
        access.require(con, username, allowed=allowed)


def account_row(con: sqlite3.Connection, account_id: int) -> sqlite3.Row:
    row = con.execute("SELECT * FROM accounts WHERE id=? AND active=1", (account_id,)).fetchone()
    if not row:
        raise HTTPException(404, "Account not found")
    return row


def account_config(con: sqlite3.Connection, account_id: int) -> dict[str, Any]:
    row = con.execute("SELECT * FROM account_config WHERE account_id=?", (account_id,)).fetchone()
    if row:
        return dict(row)
    return {
        "account_id": account_id,
        "coverage_enabled": 0,
        "coverage_period": "week",
        "audits_per_associate": 3,
        "identifier_column_default": None,
        "associate_column_default": None,
        "exclude_previously_sampled": 1,
        "case_insensitive_ids": 1,
    }


def jsonable_row(row: sqlite3.Row) -> dict[str, Any]:
    return dict(row)


class SetupIn(BaseModel):
    password: str = Field(min_length=8, max_length=200)


class LoginIn(BaseModel):
    username: str = "admin"
    password: str


class AccountIn(BaseModel):
    name: str = Field(min_length=1, max_length=120)


class ConfigIn(BaseModel):
    coverage_enabled: bool = False
    coverage_period: str = "week"
    audits_per_associate: int = Field(default=3, ge=1, le=1000)
    identifier_column_default: str | None = None
    associate_column_default: str | None = None
    exclude_previously_sampled: bool = True
    case_insensitive_ids: bool = True


class InspectIn(BaseModel):
    sheet_name: str
    header_row: int = Field(ge=1, le=200)


class FilterRule(BaseModel):
    column: str
    operator: str
    value: str


class PreviewIn(BaseModel):
    account_id: int
    process_id: int
    sheet_name: str
    header_row: int = Field(ge=1, le=200)
    identifier_column: str
    associate_column: str | None = None
    sampling_method: str = "random"
    requested_count: int = Field(default=10, ge=1, le=100_000)
    filters: list[FilterRule] = Field(default_factory=list)


class RunIn(PreviewIn):
    pass


class VoidIn(BaseModel):
    reason: str = Field(min_length=3, max_length=500)


def clean_account_name(name: str) -> str:
    name = re.sub(r"\s+", " ", name.strip())
    if not name:
        raise HTTPException(400, "Account name is required")
    return name


def safe_display_value(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        fmt = (cell.number_format or "").strip()
        if isinstance(value, (int, float)) and float(value).is_integer() and re.fullmatch(r"0+", fmt):
            return f"{int(value):0{len(fmt)}d}"
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        return str(value)
    return str(value)


def workbook_sheets(path: Path, ext: str) -> list[str]:
    if ext == ".csv":
        return ["CSV"]
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        names = wb.sheetnames
        wb.close()
        return names
    except Exception as exc:
        raise HTTPException(400, f"Workbook could not be opened: {exc}") from exc


def read_matrix(path: Path, ext: str, sheet_name: str) -> list[list[str]]:
    rows: list[list[str]] = []
    if ext == ".csv":
        try:
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as f:
                for idx, row in enumerate(csv.reader(f), start=1):
                    if idx > MAX_ROWS:
                        raise HTTPException(400, f"File exceeds the {MAX_ROWS:,}-row MVP limit")
                    rows.append([str(x) for x in row[:MAX_COLS]])
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(400, f"CSV could not be read: {exc}") from exc
        return rows
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        if sheet_name not in wb.sheetnames:
            wb.close()
            raise HTTPException(400, "Selected sheet does not exist")
        ws = wb[sheet_name]
        actual_max_col = max(1, min(ws.max_column or 1, MAX_COLS))
        for idx, cells in enumerate(ws.iter_rows(max_col=actual_max_col), start=1):
            if idx > MAX_ROWS:
                wb.close()
                raise HTTPException(400, f"Sheet exceeds the {MAX_ROWS:,}-row MVP limit")
            rows.append([safe_display_value(c) for c in cells])
        wb.close()
        return rows
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, f"Sheet could not be read: {exc}") from exc


def detect_header(matrix: list[list[str]]) -> int:
    best_idx, best_score = 0, -1
    for idx, row in enumerate(matrix[:50]):
        values = [str(x).strip() for x in row]
        nonblank = [x for x in values if x]
        if len(nonblank) < 2:
            continue
        score = len(nonblank) * 2 + len(set(nonblank))
        if score > best_score:
            best_score, best_idx = score, idx
    return best_idx + 1


def dedupe_headers(row: list[str]) -> list[str]:
    headers: list[str] = []
    seen: dict[str, int] = {}
    width = len(row)
    for i in range(width):
        base = str(row[i] or f"Column {i+1}").strip() or f"Column {i+1}"
        key = base.casefold()
        seen[key] = seen.get(key, 0) + 1
        headers.append(base if seen[key] == 1 else f"{base} ({seen[key]})")
    return headers


def matrix_to_records(matrix: list[list[str]], header_row: int) -> tuple[list[str], list[dict[str, str]]]:
    if not matrix or header_row < 1 or header_row > len(matrix):
        raise HTTPException(400, "Header row is outside the available data")
    headers = dedupe_headers(matrix[header_row - 1])
    records: list[dict[str, str]] = []
    for row_index, row in enumerate(matrix[header_row:], start=header_row + 1):
        padded = list(row) + [""] * max(0, len(headers) - len(row))
        record = {h: str(padded[i]).strip() if i < len(padded) else "" for i, h in enumerate(headers)}
        if any(v != "" for v in record.values()):
            record["__source_row__"] = str(row_index)
            records.append(record)
    return headers, records


def detect_identifier(headers: list[str], records: list[dict[str, str]]) -> dict[str, Any]:
    best = {"name": headers[0] if headers else "", "score": -999.0, "uniqueness": 0, "populated": 0}
    for header in headers:
        vals = [r.get(header, "").strip() for r in records]
        populated_vals = [v for v in vals if v]
        uniqueness = len(set(populated_vals)) / len(populated_vals) if populated_vals else 0
        populated = len(populated_vals) / len(records) if records else 0
        name = header.casefold()
        score = uniqueness * 50 + populated * 20
        if re.search(r"\bid\b|case number|ticket|reference|record|identifier|pon|order", name):
            score += 35
        if re.search(r"description|subject|time|min|date|comment", name):
            score -= 15
        if score > best["score"]:
            best = {"name": header, "score": score, "uniqueness": round(uniqueness * 100, 1), "populated": round(populated * 100, 1)}
    return best


def detect_associate(headers: list[str]) -> str | None:
    keys = ["associate", "agent", "employee", "employee id", "emp id", "empid", "advisor", "user", "owner", "executive", "auditor"]
    matches: list[tuple[int, str]] = []
    for h in headers:
        n = h.casefold()
        score = max((15 if n == k else 7 if k in n else 0) for k in keys)
        if score:
            matches.append((score, h))
    if not matches:
        return None
    matches.sort(key=lambda x: (-x[0], x[1].casefold()))
    return matches[0][1]


def upload_record(upload_id: str) -> dict[str, Any]:
    with db() as con:
        row = con.execute("SELECT * FROM uploads WHERE upload_id=? AND status='READY'", (upload_id,)).fetchone()
        if not row:
            raise HTTPException(404, "Upload not found or expired")
        return dict(row)


def normalize_identifier(value: str, case_insensitive: bool) -> str:
    value = re.sub(r"\s+", " ", str(value).strip())
    return value.casefold() if case_insensitive else value


def apply_filter(value: str, operator: str, target: str) -> bool:
    x, y = str(value or ""), str(target or "")
    op = operator.casefold().replace(" ", "_")
    if op in {"greater_than", "less_than"}:
        try:
            a, b = float(x.replace(",", "")), float(y.replace(",", ""))
            return a > b if op == "greater_than" else a < b
        except ValueError:
            return False
    xl, yl = x.casefold(), y.casefold()
    if op == "equals":
        return xl == yl
    if op == "not_equals":
        return xl != yl
    if op == "starts_with":
        return xl.startswith(yl)
    if op == "contains":
        return yl in xl
    raise HTTPException(400, f"Unsupported filter operator: {operator}")


def period_key(now: datetime, mode: str) -> str:
    d = now.astimezone().date()
    if mode == "day":
        return d.isoformat()
    if mode == "month":
        return d.strftime("%Y-%m")
    if mode != "week":
        raise HTTPException(400, "Coverage period must be day, week, or month")
    monday = d - timedelta(days=d.weekday())
    return monday.isoformat()


def build_pool(payload: PreviewIn, upload: dict[str, Any]) -> dict[str, Any]:
    with storage.materialize(upload["stored_path"], upload["file_type"]) as path:
        matrix = read_matrix(path, upload["file_type"], payload.sheet_name)
    headers, records = matrix_to_records(matrix, payload.header_row)
    if payload.identifier_column not in headers:
        raise HTTPException(400, "Identifier column is not present in the selected header row")
    if payload.associate_column and payload.associate_column not in headers:
        raise HTTPException(400, "Associate column is not present in the selected header row")
    for rule in payload.filters:
        if rule.column not in headers:
            raise HTTPException(400, f"Filter column not found: {rule.column}")

    with db() as con:
        account = account_row(con, payload.account_id)
        process = con.execute("SELECT id FROM processes WHERE id=? AND account_id=? AND active=1", (payload.process_id, payload.account_id)).fetchone()
        if not process:
            raise HTTPException(400, "Select an active process belonging to this account")
        cfg = sampling_config(con, payload.process_id)
        case_insensitive = bool(cfg["case_insensitive_ids"])
        exclude_previous = bool(cfg["exclude_previously_sampled"])
        previous = set()
        if exclude_previous:
            previous = {
                normalize_identifier(r[0], case_insensitive)
                for r in con.execute(
                    """
                    SELECT sr.identifier_raw
                    FROM sample_records sr
                    JOIN sampling_runs rr ON rr.run_id=sr.run_id
                    WHERE rr.process_id=? AND rr.status='COMPLETED'
                    """,
                    (payload.process_id,),
                ).fetchall()
            }

    blank_ids = 0
    duplicate_ids = 0
    previous_count = 0
    filter_excluded = 0
    seen: set[str] = set()
    eligible: list[dict[str, Any]] = []
    coverage_roster: set[str] = set()
    blank_associate_candidates = 0
    for r in records:
        raw_id = r.get(payload.identifier_column, "").strip()
        if not raw_id:
            blank_ids += 1
            continue
        norm = normalize_identifier(raw_id, case_insensitive)
        if norm in seen:
            duplicate_ids += 1
            continue
        seen.add(norm)
        passes = all(apply_filter(r.get(rule.column, ""), rule.operator, rule.value) for rule in payload.filters)
        if not passes:
            filter_excluded += 1
            continue
        if payload.associate_column:
            candidate_associate = r.get(payload.associate_column, "").strip()
            if candidate_associate:
                coverage_roster.add(candidate_associate)
            else:
                blank_associate_candidates += 1
        if exclude_previous and norm in previous:
            previous_count += 1
            continue
        eligible.append({"row": r, "identifier_raw": raw_id, "normalized_identifier": norm})

    total_rows = len(records)
    valid_unique = len(seen)
    coverage = None
    cfg_enabled = bool(cfg["coverage_enabled"])
    if cfg_enabled:
        if not payload.associate_column:
            raise HTTPException(400, "This process enforces coverage sampling; confirm the associate column")
        pk = period_key(utcnow(), cfg["coverage_period"])
        with db() as con:
            completed_by_associate = {
                row["associate"] or "": row["n"]
                for row in con.execute(
                    """
                    SELECT COALESCE(sr.associate,'') associate, COUNT(*) n
                    FROM sample_records sr
                    JOIN sampling_runs rr ON rr.run_id=sr.run_id
                    WHERE rr.process_id=? AND rr.status='COMPLETED' AND sr.period_key=?
                    GROUP BY COALESCE(sr.associate,'')
                    """,
                    (payload.process_id, pk),
                ).fetchall()
            }
        groups: dict[str, list[dict[str, Any]]] = {}
        for item in eligible:
            associate = item["row"].get(payload.associate_column, "").strip()
            if not associate:
                continue
            groups.setdefault(associate, []).append(item)
        quota = int(cfg["audits_per_associate"])
        details = []
        total_required = total_feasible = 0
        for associate in sorted(coverage_roster, key=str.casefold):
            completed = int(completed_by_associate.get(associate, 0))
            remaining = max(0, quota - completed)
            available = len(groups.get(associate, []))
            feasible = min(remaining, available)
            shortfall = max(0, remaining - available)
            total_required += remaining
            total_feasible += feasible
            details.append({
                "associate": associate,
                "required": quota,
                "completed": completed,
                "remaining": remaining,
                "eligible": available,
                "shortfall": shortfall,
            })
        coverage = {
            "period": cfg["coverage_period"],
            "period_key": pk,
            "quota": quota,
            "blank_associate_records": blank_associate_candidates,
            "total_required": total_required,
            "total_feasible": total_feasible,
            "associates_with_shortfall": sum(1 for d in details if d["shortfall"] > 0),
            "details": details,
            "groups": groups,
        }

    return {
        "account": account["name"],
        "headers": headers,
        "records": records,
        "eligible": eligible,
        "config": cfg,
        "coverage": coverage,
        "validation": {
            "source_rows": total_rows,
            "valid_unique_identifiers": valid_unique,
            "blank_identifiers": blank_ids,
            "duplicate_identifiers": duplicate_ids,
            "previously_sampled": previous_count,
            "filter_excluded": filter_excluded,
            "eligible_population": len(eligible),
        },
    }


def run_id() -> str:
    return f"QSR-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{secrets.token_hex(2).upper()}"


def derived_seed(seed: int, key: str) -> int:
    digest = hashlib.sha256(f"{seed}|{key}".encode()).digest()
    return int.from_bytes(digest[:8], "big")


def select_sample(pool_data: dict[str, Any], payload: RunIn, seed: int) -> list[dict[str, Any]]:
    coverage = pool_data["coverage"]
    if coverage:
        selected: list[dict[str, Any]] = []
        quota = coverage["quota"]
        completed = {d["associate"]: d["completed"] for d in coverage["details"]}
        for associate in sorted(coverage["groups"], key=str.casefold):
            gap = max(0, quota - completed.get(associate, 0))
            group = sorted(
                coverage["groups"][associate],
                key=lambda x: (x["normalized_identifier"], int(x["row"].get("__source_row__", "0") or 0)),
            )
            rng = random.Random(derived_seed(seed, associate))
            rng.shuffle(group)
            selected.extend(group[:gap])
        selected.sort(key=lambda x: (x["row"].get(payload.associate_column or "", "").casefold(), x["normalized_identifier"]))
        return selected
    pool = sorted(
        pool_data["eligible"],
        key=lambda x: (x["normalized_identifier"], int(x["row"].get("__source_row__", "0") or 0)),
    )
    if payload.requested_count > len(pool):
        raise HTTPException(400, f"Only {len(pool)} eligible unique records are available")
    rng = random.Random(seed)
    rng.shuffle(pool)
    return pool[: payload.requested_count]


def formula_safe(value: Any) -> Any:
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def backup_database(prefix: str = "Manual") -> Path:
    name = f"{prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
    target = BACKUP_DIR / name
    src = connect()
    dst = sqlite3.connect(target)
    try:
        src.backup(dst)
    finally:
        dst.close()
        src.close()
    return target


def prune_auto_backups(keep: int = 7) -> None:
    files = sorted(BACKUP_DIR.glob("Auto_*.db"), key=lambda p: p.stat().st_mtime, reverse=True)
    for old in files[keep:]:
        old.unlink(missing_ok=True)


def ensure_daily_auto_backup() -> None:
    if uses_postgres():
        return
    with db() as con:
        initialized = con.execute("SELECT COUNT(*) FROM users").fetchone()[0] > 0
    if not initialized:
        return
    today = datetime.now().strftime("%Y%m%d")
    if any(BACKUP_DIR.glob(f"Auto_{today}_*.db")):
        return
    target = backup_database("Auto")
    prune_auto_backups(7)
    logger.info("automatic_backup=%s", target.name)


def cleanup_stale_uploads(hours: int = 24) -> None:
    cutoff = (utcnow() - timedelta(hours=hours)).replace(microsecond=0).isoformat()
    with db() as con:
        rows = con.execute("SELECT upload_id, stored_path FROM uploads WHERE status='READY' AND created_at < ?", (cutoff,)).fetchall()
        for row in rows:
            storage.delete(row["stored_path"])
            con.execute("UPDATE uploads SET status='EXPIRED' WHERE upload_id=?", (row["upload_id"],))
    if rows:
        logger.info("expired_uploads=%s", len(rows))


app = FastAPI(title=APP_NAME, version=APP_VERSION, docs_url=None, redoc_url=None)
app.mount("/static", StaticFiles(directory=BASE_DIR / "static"), name="static")


@app.get("/healthz", include_in_schema=False)
def healthz():
    return {"status": "ok"}


@app.get("/readyz", include_in_schema=False)
async def readyz():
    if uses_postgres():
        try:
            with db() as con:
                for statement in (
                    "SELECT username,account_id,role_name FROM account_user_roles LIMIT 0",
                    "SELECT username,role_name FROM legacy_user_roles LIMIT 0",
                    "SELECT process_id,coverage_enabled FROM process_sampling_config LIMIT 0",
                    "SELECT created_by FROM uploads LIMIT 0",
                ):
                    con.execute(statement)
        except Exception:
            logger.error("readiness_failed: PostgreSQL account schema unavailable; run alembic upgrade head using the migration role")
            return JSONResponse(status_code=503, content={"status": "not_ready", "detail": "Database schema unavailable. Run alembic upgrade head using the migration role."})
    return {"status": "ready"}


@app.on_event("startup")
def startup_backup():
    try:
        cleanup_stale_uploads()
        ensure_daily_auto_backup()
    except Exception:
        logger.exception("Startup maintenance failed")


@app.middleware("http")
async def security_headers(request: Request, call_next):
    if request.method not in {"GET", "HEAD", "OPTIONS"} and request.url.path not in {"/api/setup", "/api/login"}:
        session_cookie = request.cookies.get("qsr_session")
        if session_cookie:
            csrf_cookie = request.cookies.get("qsr_csrf")
            csrf_header = request.headers.get("X-CSRF-Token")
            if not csrf_cookie or not csrf_header or not hmac.compare_digest(csrf_cookie, csrf_header):
                return JSONResponse(status_code=403, content={"detail": "CSRF validation failed"})
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    response.headers["Content-Security-Policy"] = "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'"
    return response


@app.get("/")
def home():
    return FileResponse(BASE_DIR / "static" / "index.html")


@app.get("/api/status")
def status():
    with db() as con:
        initialized = con.execute("SELECT COUNT(*) FROM users").fetchone()[0] > 0
    return {"app": APP_NAME, "version": APP_VERSION, "initialized": initialized}


@app.post("/api/setup")
def setup(payload: SetupIn, response: Response):
    with db() as con:
        if con.execute("SELECT COUNT(*) FROM users").fetchone()[0] > 0:
            raise HTTPException(409, "Application is already initialized")
        salt, digest = password_hash(payload.password)
        con.execute(
            "INSERT INTO users(username, password_salt, password_hash, created_at) VALUES(?,?,?,?)",
            ("admin", salt, digest, iso_now()),
        )
        con.execute("INSERT INTO user_profiles(username,display_name,active,must_change_password,updated_at) VALUES('admin','Administrator',1,0,?)", (iso_now(),))
        con.execute("INSERT INTO user_roles(username,role_name) VALUES('admin','Administrator')")
    token = create_session("admin")
    csrf = secrets.token_urlsafe(24)
    set_auth_cookies(response, token, csrf)
    audit("SETUP_COMPLETE", "admin", "application", APP_VERSION)
    return {"ok": True, "username": "admin", "csrf_token": csrf}


@app.post("/api/login")
def login(payload: LoginIn, response: Response):
    with db() as con:
        row = con.execute("""SELECT u.*, p.active FROM users u
                             LEFT JOIN user_profiles p ON p.username=u.username
                             WHERE u.username=?""", (payload.username.strip(),)).fetchone()
    if not row or not row["active"] or not verify_password(payload.password, row["password_salt"], row["password_hash"]):
        logger.warning("login_failed user=%s", payload.username)
        raise HTTPException(401, "Invalid username or password")
    token = create_session(row["username"], expected_password_hash=row["password_hash"])
    csrf = secrets.token_urlsafe(24)
    set_auth_cookies(response, token, csrf)
    audit("LOGIN", row["username"], "session", None)
    return {"ok": True, "username": row["username"], "csrf_token": csrf}


def set_auth_cookies(response: Response, token: str, csrf: str) -> None:
    secure = settings.is_production_like or settings.public_origin.lower().startswith("https://")
    response.set_cookie("qsr_session", token, httponly=True, samesite="strict", secure=secure, max_age=SESSION_HOURS * 3600)
    response.set_cookie("qsr_csrf", csrf, httponly=False, samesite="strict", secure=secure, max_age=SESSION_HOURS * 3600)


@app.post("/api/logout")
def logout(response: Response, qsr_session: str | None = Cookie(default=None), user: str = Depends(current_user)):
    if qsr_session:
        with db() as con:
            con.execute("DELETE FROM sessions WHERE token_hash=?", (hashlib.sha256(qsr_session.encode()).hexdigest(),))
    response.delete_cookie("qsr_session")
    response.delete_cookie("qsr_csrf")
    audit("LOGOUT", user, "session", None)
    return {"ok": True}


@app.get("/api/me")
def me(user: str = Depends(current_user)):
    return {"username": user}


@app.get("/api/dashboard")
def dashboard(user: str = Depends(current_user)):
    with db() as con:
        clause, params = access.scope(con, user, "a.id")
        summary = [dict(r) for r in con.execute(f"""SELECT a.id account_id,a.name account,
            SUM(CASE WHEN r.status='COMPLETED' THEN 1 ELSE 0 END) runs,
            COALESCE(SUM(CASE WHEN r.status='COMPLETED' THEN r.selected_count ELSE 0 END),0) samples,
            SUM(CASE WHEN r.status='VOIDED' THEN 1 ELSE 0 END) voided_runs,
            MAX(CASE WHEN r.status='COMPLETED' THEN r.created_at END) last_run
            FROM accounts a LEFT JOIN sampling_runs r ON r.account_id=a.id
            WHERE a.active=1 AND {clause} GROUP BY a.id,a.name ORDER BY a.name""", params).fetchall()]
    return {"metrics": {"accounts": len(summary), "samples": sum(r["samples"] for r in summary),
        "runs": sum(r["runs"] for r in summary), "voided_runs": sum(r["voided_runs"] for r in summary)}, "summary": summary}


@app.get("/api/accounts")
def list_accounts(user: str = Depends(current_user), include_archived: bool = False):
    with db() as con:
        clause, params = access.scope(con, user, "id")
        active = "1=1" if include_archived else "active=1"
        rows = [dict(r) for r in con.execute(f"SELECT * FROM accounts WHERE {active} AND {clause} ORDER BY name COLLATE NOCASE", params).fetchall()]
        for row in rows:
            row["config"] = account_config(con, row["id"])
    return rows


@app.post("/api/accounts")
def create_account(payload: AccountIn, user: str = Depends(current_user)):
    ensure_roles(user, "Administrator")
    name = clean_account_name(payload.name)
    now = iso_now()
    try:
        with db() as con:
            cur = con.execute("INSERT INTO accounts(name, active, created_at, updated_at) VALUES(?,?,?,?)", (name, 1, now, now))
            account_id = cur.lastrowid
            con.execute("INSERT INTO account_config(account_id, updated_at) VALUES(?,?)", (account_id, now))
            pcur = con.execute(
                "INSERT INTO processes(account_id,name,process_type,created_at,updated_at) VALUES(?,?,?,?,?)",
                (account_id, "General Service Process", "back_office", now, now),
            )
            con.execute("INSERT INTO process_settings(process_id,updated_at) VALUES(?,?)", (pcur.lastrowid, now))
            initialize_process(con, pcur.lastrowid, now)
    except sqlite3.IntegrityError:
        raise HTTPException(409, "Account already exists")
    audit("ACCOUNT_CREATED", user, "account", str(account_id), {"name": name})
    return {"id": account_id, "name": name}


@app.delete("/api/accounts/{account_id}")
def archive_account(account_id: int, user: str = Depends(current_user)):
    ensure_roles(user, "Administrator")
    with db() as con:
        row = account_row(con, account_id)
        con.execute("UPDATE accounts SET active=0, updated_at=? WHERE id=?", (iso_now(), account_id))
    audit("ACCOUNT_ARCHIVED", user, "account", str(account_id), {"name": row["name"]})
    return {"ok": True}


@app.get("/api/accounts/{account_id}/config")
def get_config(account_id: int, user: str = Depends(current_user)):
    with db() as con:
        access.require(con, user, account_id)
        account_row(con, account_id)
        return account_config(con, account_id)


@app.put("/api/accounts/{account_id}/config")
def save_config(account_id: int, payload: ConfigIn, user: str = Depends(current_user)):
    ensure_roles(user, "Administrator")
    raise HTTPException(410, "Sampling controls are now per process. Use /api/admin/processes/{process_id}/sampling-controls")


@app.post("/api/uploads")
async def upload_workbook(file: UploadFile = File(...), user: str = Depends(current_user)):
    ensure_roles(user, "Administrator", "QA Auditor")
    filename = Path(file.filename or "upload").name
    ext = Path(filename).suffix.casefold()
    if ext not in {".xlsx", ".csv"}:
        raise HTTPException(400, "Production MVP supports XLSX and CSV files only")
    upload_id = uuid.uuid4().hex
    target = UPLOAD_DIR / f"{upload_id}{ext}"
    size = 0
    sha = hashlib.sha256()
    try:
        with target.open("wb") as out:
            while True:
                chunk = await file.read(1024 * 1024)
                if not chunk:
                    break
                size += len(chunk)
                if size > MAX_FILE_BYTES:
                    raise HTTPException(400, f"File exceeds the {MAX_FILE_BYTES // (1024*1024)} MB limit")
                sha.update(chunk)
                out.write(chunk)
        sheets = workbook_sheets(target, ext)
        matrix = read_matrix(target, ext, sheets[0])
        if not matrix or not any(any(str(v).strip() for v in row) for row in matrix):
            raise HTTPException(400, "File contains no usable data")
        detected_header = detect_header(matrix)
        stored_reference = storage.put_file(target, storage.object_key(upload_id, ext))
        with db() as con:
            con.execute(
                "INSERT INTO uploads(upload_id, original_name, stored_path, file_type, sha256, size_bytes, created_at, created_by, status) VALUES(?,?,?,?,?,?,?,?,'READY')",
                (upload_id, filename, stored_reference, ext, sha.hexdigest(), size, iso_now(), user),
            )
        if storage.remote:
            target.unlink(missing_ok=True)
        audit("FILE_UPLOADED", user, "upload", upload_id, {"filename": filename, "sha256": sha.hexdigest(), "size_bytes": size})
        return {"upload_id": upload_id, "filename": filename, "sha256": sha.hexdigest(), "size_bytes": size, "sheets": sheets, "detected_header_row": detected_header}
    except Exception:
        target.unlink(missing_ok=True)
        raise


@app.post("/api/uploads/{upload_id}/inspect")
def inspect_upload(upload_id: str, payload: InspectIn, user: str = Depends(current_user)):
    ensure_roles(user, "Administrator", "QA Auditor")
    with db() as con:
        access.staging(con, user, "uploads", "upload_id", upload_id)
    upload = upload_record(upload_id)
    with storage.materialize(upload["stored_path"], upload["file_type"]) as path:
        matrix = read_matrix(path, upload["file_type"], payload.sheet_name)
    headers, records = matrix_to_records(matrix, payload.header_row)
    if not headers or not records:
        raise HTTPException(400, "Selected header row does not produce usable records")
    identifier = detect_identifier(headers, records)
    associate = detect_associate(headers)
    return {
        "columns": headers,
        "identifier_suggestion": identifier,
        "associate_suggestion": associate,
        "source_rows": len(records),
        "preview_rows": [{k: v for k, v in r.items() if k != "__source_row__"} for r in records[:8]],
    }


@app.post("/api/uploads/{upload_id}/preview")
def preview_sample(upload_id: str, payload: PreviewIn, user: str = Depends(current_user)):
    ensure_roles(user, "Administrator", "QA Auditor")
    with db() as con:
        access.staging(con, user, "uploads", "upload_id", upload_id)
        access.process(con, user, payload.process_id, ("QA Auditor",), active=True, account_id=payload.account_id)
    upload = upload_record(upload_id)
    pool_data = build_pool(payload, upload)
    coverage = pool_data["coverage"]
    out = {
        "account": pool_data["account"],
        "validation": pool_data["validation"],
        "coverage_enforced": bool(pool_data["config"]["coverage_enabled"]),
        "coverage": None,
        "requested_count": payload.requested_count,
        "can_generate": True,
    }
    if coverage:
        out["coverage"] = {k: v for k, v in coverage.items() if k != "groups"}
        out["requested_count"] = coverage["total_required"]
        out["feasible_count"] = coverage["total_feasible"]
        out["can_generate"] = coverage["total_feasible"] > 0
    else:
        out["feasible_count"] = min(payload.requested_count, len(pool_data["eligible"]))
        out["can_generate"] = payload.requested_count <= len(pool_data["eligible"]) and len(pool_data["eligible"]) > 0
    return out


@app.post("/api/uploads/{upload_id}/runs")
def generate_run(upload_id: str, payload: RunIn, user: str = Depends(current_user)):
    ensure_roles(user, "Administrator", "QA Auditor")
    with db() as con:
        access.staging(con, user, "uploads", "upload_id", upload_id)
        access.process(con, user, payload.process_id, ("QA Auditor",), active=True, account_id=payload.account_id)
    upload = upload_record(upload_id)
    pool_data = build_pool(payload, upload)
    seed = secrets.randbits(63)
    selected = select_sample(pool_data, payload, seed)
    if not selected:
        raise HTTPException(400, "No records are available for selection")
    rid = run_id()
    now = iso_now()
    cfg = pool_data["config"]
    coverage = pool_data["coverage"]
    requested = coverage["total_required"] if coverage else payload.requested_count
    pk = coverage["period_key"] if coverage else None
    method = "coverage" if coverage else payload.sampling_method
    with db() as con:
        con.execute("UPDATE processes SET updated_at=updated_at WHERE id=?", (payload.process_id,))
        access.process(con, user, payload.process_id, ("QA Auditor",), active=True, account_id=payload.account_id)
        initialize_process(con, payload.process_id, now)
        account = account_row(con, payload.account_id)
        con.execute(
            """
            INSERT INTO sampling_runs(run_id, account_id, process_id, upload_id, source_filename, source_hash, sheet_name, header_row,
                identifier_column, associate_column, sampling_method, requested_count, eligible_count, selected_count, seed,
                algorithm_version, coverage_enabled, coverage_period, coverage_quota, period_key, filters_json, validation_json,
                coverage_summary_json, app_version, created_at, created_by, status)
            VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?, 'COMPLETED')
            """,
            (
                rid, payload.account_id, payload.process_id, upload_id, upload["original_name"], upload["sha256"], payload.sheet_name, payload.header_row,
                payload.identifier_column, payload.associate_column, method, requested, pool_data["validation"]["eligible_population"],
                len(selected), str(seed), ALGORITHM_VERSION, int(bool(coverage)), cfg["coverage_period"] if coverage else None,
                cfg["audits_per_associate"] if coverage else None, pk, json.dumps([r.model_dump() for r in payload.filters], ensure_ascii=False),
                json.dumps(pool_data["validation"], ensure_ascii=False),
                json.dumps({k: v for k, v in coverage.items() if k != "groups"} if coverage else {}, ensure_ascii=False),
                APP_VERSION, now, user,
            ),
        )
        for seq, item in enumerate(selected, start=1):
            row = {k: v for k, v in item["row"].items() if k != "__source_row__"}
            associate = row.get(payload.associate_column, "").strip() if payload.associate_column else ""
            rec_cur = con.execute(
                """
                INSERT INTO sample_records(run_id, identifier_raw, normalized_identifier, associate, selection_sequence,
                    row_json, period_key, selected_at) VALUES(?,?,?,?,?,?,?,?)
                """,
                (rid, item["identifier_raw"], item["normalized_identifier"], associate, seq, json.dumps(row, ensure_ascii=False), pk, now),
            )
            scorecard = con.execute(
                "SELECT id FROM scorecard_versions WHERE process_id=? AND status='PUBLISHED' ORDER BY version DESC LIMIT 1",
                (payload.process_id,),
            ).fetchone()
            audit_id = f"AUD-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{secrets.token_hex(3).upper()}"
            con.execute(
                """INSERT INTO audit_cases(audit_id,process_id,sample_record_id,external_case_id,scorecard_version_id,
                   status,associate,source_json,created_at,created_by) VALUES(?,?,?,?,?,'UNASSIGNED',?,?,?,?)""",
                (audit_id, payload.process_id, rec_cur.lastrowid, f"{rid}:{item['identifier_raw']}",
                 scorecard["id"] if scorecard else None, associate, json.dumps(row, ensure_ascii=False), now, user),
            )
        # Save successful mappings for this process only.
        con.execute(
            """
            UPDATE process_sampling_config SET identifier_column_default=?, associate_column_default=?, updated_at=? WHERE process_id=?
            """,
            (payload.identifier_column, payload.associate_column, now, payload.process_id),
        )
        con.execute("UPDATE uploads SET status='USED' WHERE upload_id=?", (upload_id,))
    storage.delete(upload["stored_path"])
    audit("RUN_COMPLETED", user, "run", rid, {"account": account["name"], "selected_count": len(selected), "seed": str(seed)})
    return {"run_id": rid, "selected_count": len(selected), "status": "COMPLETED"}


@app.get("/api/runs")
def list_runs(account_id: int | None = None, limit: int = 200, user: str = Depends(current_user), process_id: int | None = None):
    with db() as con:
        clause, params = access.scope(con, user, "r.account_id", account_id, process_id)
        if account_id is not None:
            clause += " AND r.account_id=?"; params.append(account_id)
        if process_id is not None:
            clause += " AND r.process_id=?"; params.append(process_id)
        params.append(max(1, min(limit, 500)))
        rows = con.execute(f"""SELECT r.*,a.name account_name,p.name process_name,p.active process_active
            FROM sampling_runs r JOIN accounts a ON a.id=r.account_id JOIN processes p ON p.id=r.process_id
            WHERE {clause} ORDER BY r.created_at DESC LIMIT ?""", params).fetchall()
    return [dict(r) for r in rows]


@app.get("/api/runs/{rid}")
def get_run(rid: str, user: str = Depends(current_user)):
    with db() as con:
        access.record(con, user, "sampling_runs", "run_id", rid, ())
        run = con.execute("SELECT r.*,a.name account_name FROM sampling_runs r JOIN accounts a ON a.id=r.account_id WHERE run_id=?", (rid,)).fetchone()
        if not run:
            raise HTTPException(404, "Run not found")
        records = [dict(r) for r in con.execute("SELECT * FROM sample_records WHERE run_id=? ORDER BY selection_sequence", (rid,)).fetchall()]
    out = dict(run)
    out["filters"] = json.loads(out.pop("filters_json") or "[]")
    out["validation"] = json.loads(out.pop("validation_json", "{}") or "{}")
    out["coverage_summary"] = json.loads(out.pop("coverage_summary_json", "{}") or "{}")
    parsed_records = []
    for rec in records:
        rec_out = dict(rec)
        rec_out["row"] = json.loads(rec_out.pop("row_json"))
        parsed_records.append(rec_out)
    out["records"] = parsed_records
    return out


@app.post("/api/runs/{rid}/void")
def void_run(rid: str, payload: VoidIn, user: str = Depends(current_user)):
    ensure_roles(user, "Administrator", "QA Reviewer")
    with db() as con:
        access.record(con, user, "sampling_runs", "run_id", rid, ('QA Reviewer',))
        row = con.execute("SELECT status FROM sampling_runs WHERE run_id=?", (rid,)).fetchone()
        if not row:
            raise HTTPException(404, "Run not found")
        if row["status"] == "VOIDED":
            raise HTTPException(409, "Run is already voided")
        con.execute("UPDATE sampling_runs SET status='VOIDED', void_reason=?, voided_at=?, voided_by=? WHERE run_id=?", (payload.reason.strip(), iso_now(), user, rid))
        con.execute("UPDATE audit_cases SET status='VOIDED',void_reason=? WHERE sample_record_id IN (SELECT id FROM sample_records WHERE run_id=?)", (payload.reason.strip(), rid))
    audit("RUN_VOIDED", user, "run", rid, {"reason": payload.reason.strip()})
    return {"ok": True}


@app.get("/api/inventory")
def inventory(account_id: int | None = None, search: str = "", limit: int = 500, user: str = Depends(current_user)):
    limit = max(1, min(limit, 1000))
    params: list[Any] = []
    where = ["r.status='COMPLETED'"]
    if account_id:
        where.append("r.account_id=?")
        params.append(account_id)
    if search.strip():
        where.append("(sr.identifier_raw LIKE ? OR sr.associate LIKE ? OR r.run_id LIKE ?)")
        q = f"%{search.strip()}%"
        params.extend([q, q, q])
    params.append(limit)
    with db() as con:
        scoped, scope_params = access.scope(con, user, "r.account_id", account_id)
        where.append(scoped)
        params[-1:-1] = scope_params
        rows = con.execute(
            f"""
            SELECT sr.identifier_raw identifier, sr.associate, sr.period_key, sr.selection_sequence,
                   r.run_id, r.created_at selected_at, r.sampling_method, a.name account
            FROM sample_records sr JOIN sampling_runs r ON r.run_id=sr.run_id JOIN accounts a ON a.id=r.account_id
            WHERE {' AND '.join(where)} ORDER BY r.created_at DESC, sr.selection_sequence LIMIT ?
            """,
            params,
        ).fetchall()
    return [dict(r) for r in rows]


@app.get("/api/runs/{rid}/export")
def export_run(rid: str, user: str = Depends(current_user)):
    detail = get_run(rid, user)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sample Results"
    records = detail["records"]
    source_headers: list[str] = []
    if records:
        source_headers = list(records[0]["row"].keys())
    headers = source_headers + ["Run ID", "Account", "Sampling Method", "Selection Sequence", "Selected At"]
    ws.append([formula_safe(value) for value in headers])
    for rec in records:
        values = [formula_safe(rec["row"].get(h, "")) for h in source_headers]
        values += [detail["run_id"], detail["account_name"], detail["sampling_method"], rec["selection_sequence"], rec["selected_at"]]
        ws.append([formula_safe(value) for value in values])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    manifest = wb.create_sheet("Run Manifest")
    manifest.append(["Field", "Value"])
    manifest_fields = [
        ("Run ID", detail["run_id"]), ("Status", detail["status"]), ("Created At (UTC)", detail["created_at"]),
        ("Created By", detail["created_by"]), ("Account", detail["account_name"]), ("Source File", detail["source_filename"]),
        ("Source SHA-256", detail["source_hash"]), ("Sheet", detail["sheet_name"]), ("Header Row", detail["header_row"]),
        ("Identifier Column", detail["identifier_column"]), ("Associate Column", detail["associate_column"] or ""),
        ("Sampling Method", detail["sampling_method"]), ("Requested Count", detail["requested_count"]),
        ("Eligible Population", detail["eligible_count"]), ("Selected Count", detail["selected_count"]),
        ("Application Version", detail.get("app_version") or APP_VERSION), ("Algorithm", detail["algorithm_version"]), ("Seed", detail["seed"]),
        ("Coverage Enabled", bool(detail["coverage_enabled"])), ("Coverage Period", detail["coverage_period"] or ""),
        ("Coverage Quota", detail["coverage_quota"] or ""), ("Period Key", detail["period_key"] or ""),
        ("Source Rows", detail.get("validation", {}).get("source_rows", "")),
        ("Valid Unique Identifiers", detail.get("validation", {}).get("valid_unique_identifiers", "")),
        ("Blank Identifiers", detail.get("validation", {}).get("blank_identifiers", "")),
        ("Duplicate Identifiers", detail.get("validation", {}).get("duplicate_identifiers", "")),
        ("Previously Sampled Excluded", detail.get("validation", {}).get("previously_sampled", "")),
        ("Filtered Out", detail.get("validation", {}).get("filter_excluded", "")),
        ("Filters", json.dumps(detail["filters"], ensure_ascii=False)),
        ("Coverage Summary", json.dumps(detail.get("coverage_summary", {}), ensure_ascii=False)),
        ("Void Reason", detail["void_reason"] or ""),
    ]
    for field, value in manifest_fields:
        manifest.append([field, formula_safe(str(value)) if value is not None else ""])
    manifest.column_dimensions["A"].width = 28
    manifest.column_dimensions["B"].width = 80

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    audit("RUN_EXPORTED", user, "run", rid)
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{rid}.xlsx"'},
    )


@app.get("/api/settings/backups")
def list_backups(user: str = Depends(current_user)):
    items = []
    for path in sorted(BACKUP_DIR.glob("*.db"), key=lambda p: p.stat().st_mtime, reverse=True):
        items.append({"name": path.name, "size_bytes": path.stat().st_size, "modified_at": datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()})
    return items[:50]


@app.post("/api/settings/backup")
def create_backup(user: str = Depends(current_user)):
    ensure_roles(user, "Administrator")
    target = backup_database("Manual")
    audit("BACKUP_CREATED", user, "backup", target.name)
    return {"ok": True, "name": target.name}


@app.post("/api/settings/restore")
async def restore_backup(file: UploadFile = File(...), user: str = Depends(current_user)):
    ensure_roles(user, "Administrator")
    if Path(file.filename or "").suffix.casefold() != ".db":
        raise HTTPException(400, "Restore accepts a .db SQLite backup only")
    fd, temp_name = tempfile.mkstemp(prefix="qsr_restore_", suffix=".db", dir=DATA_DIR)
    os.close(fd)
    temp = Path(temp_name)
    try:
        with temp.open("wb") as out:
            while chunk := await file.read(1024 * 1024):
                out.write(chunk)
        test = sqlite3.connect(temp)
        try:
            integrity = test.execute("PRAGMA integrity_check").fetchone()[0]
            tables = {r[0] for r in test.execute("SELECT name FROM sqlite_master WHERE type='table'").fetchall()}
        finally:
            test.close()
        if integrity != "ok" or not REQUIRED_TABLES.issubset(tables):
            raise HTTPException(400, "Backup failed integrity/schema validation")
        pre = BACKUP_DIR / f"PreRestore_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
        src = connect(); dst = sqlite3.connect(pre)
        try:
            src.backup(dst)
        finally:
            dst.close(); src.close()
        restore_src = sqlite3.connect(temp)
        restore_dst = connect()
        try:
            restore_src.backup(restore_dst)
            restore_dst.commit()
        finally:
            restore_dst.close(); restore_src.close()
        init_quality_db()
        audit("DATABASE_RESTORED", user, "backup", file.filename, {"pre_restore_backup": pre.name})
        return {"ok": True, "pre_restore_backup": pre.name, "message": "Database restored. Sign in again if your session no longer exists."}
    finally:
        temp.unlink(missing_ok=True)


@app.get("/api/audit")
def audit_log(limit: int = 200, user: str = Depends(current_user)):
    ensure_roles(user, "Administrator")
    limit = max(1, min(limit, 500))
    with db() as con:
        rows = con.execute("SELECT * FROM audit_events ORDER BY created_at DESC LIMIT ?", (limit,)).fetchall()
    return [dict(r) for r in rows]


from quality import init_quality_db, router as quality_router

init_quality_db()
app.include_router(quality_router)


@app.exception_handler(HTTPException)
async def http_error(request: Request, exc: HTTPException):
    return JSONResponse(status_code=exc.status_code, content={"detail": exc.detail})
