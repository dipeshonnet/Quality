"""Regression coverage for bounded analytics queries and streaming imports."""
import io
import sqlite3
import unittest
from unittest.mock import patch

from fastapi import UploadFile
from openpyxl import Workbook

import test_quality as baseline
from test_quality import app, quality


class RefactoringTests(unittest.TestCase):
    setUp = baseline.QualityCommandCenterTests.setUp
    seed_process = baseline.QualityCommandCenterTests.seed_process

    def test_analytics_aggregates_without_per_audit_parameters(self):
        aid, pid = self.seed_process()
        stamp = "2026-10-08T12:00:00+00:00"
        with app.db() as con:
            for index in range(120):
                audit_id = f"AGG-{index}"
                con.execute(
                    """INSERT INTO audit_cases(audit_id,process_id,status,defect_count,
                       opportunities,critical_fail,created_at,created_by,reviewed_at)
                       VALUES(?,?,'REVIEWED',2,0,1,?,'admin',?)""",
                    (audit_id, pid, stamp, stamp),
                )
                for category in ("Accuracy", ""):
                    con.execute(
                        """INSERT INTO audit_defects(audit_id,category,name,severity,created_at)
                           VALUES(?,?, 'Missing field','Major',?)""", (audit_id, category, stamp),
                    )
            # Reproduces hosts with restrictive bind limits. A parameter per audit fails.
            con.setlimit(sqlite3.SQLITE_LIMIT_VARIABLE_NUMBER, 20)
            with patch.object(quality, "db") as database:
                database.return_value.__enter__.return_value = con
                result = quality.analytics_payload(aid, pid, stamp[:10], stamp[:10])
        self.assertEqual(result["metrics"]["units"], 120)
        self.assertEqual(result["metrics"]["defects"], 240)
        self.assertEqual(result["metrics"]["opportunities"], 120)
        self.assertEqual(result["metrics"]["critical_defects"], 120)
        self.assertEqual(len(result["control_chart"]), 1)
        self.assertEqual(result["control_chart"][0]["u"], 2)
        self.assertEqual([r["count"] for r in result["pareto"]], [120, 120])
        self.assertEqual([r["category"] for r in result["pareto"]], ["Accuracy", "Missing field"])
        self.assertEqual(result["pareto"][-1]["cumulative_percent"], 100)

    def test_empty_analytics_preserves_response(self):
        aid, pid = self.seed_process()
        result = quality.analytics_payload(aid, pid, None, None)
        self.assertEqual(result["metrics"]["units"], 0)
        self.assertEqual(result["metrics"]["yield"], 0)
        self.assertEqual(result["metrics"]["sigma"], 6)
        self.assertEqual(result["control_chart"], [])
        self.assertEqual(result["pareto"], [])
        self.assertEqual(result["stability"], "insufficient")

    def test_xlsx_import_preserves_headers_and_values(self):
        workbook = Workbook()
        workbook.active.append(["ID", None, "Count"])
        workbook.active.append([" A1 ", " Alice ", 0])
        workbook.active.append(["A2", None, 2])
        content = io.BytesIO()
        workbook.save(content)
        workbook.close()
        headers, rows = quality._read_result_file(UploadFile(filename="results.xlsx", file=io.BytesIO()), content.getvalue())
        self.assertEqual(headers, ["ID", "Column 2", "Count"])
        self.assertEqual(rows, [
            {"ID": "A1", "Column 2": "Alice", "Count": "0"},
            {"ID": "A2", "Column 2": "", "Count": "2"},
        ])

    def test_xlsx_workbook_closes_on_conversion_failure(self):
        with patch.object(quality, "load_workbook") as loader:
            workbook = loader.return_value
            workbook.active.iter_rows.side_effect = ValueError("Invalid worksheet")
            with self.assertRaises(ValueError):
                quality._read_result_file(UploadFile(filename="results.xlsx", file=io.BytesIO()), b"")
            workbook.close.assert_called_once()
